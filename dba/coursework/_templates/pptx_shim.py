#!/usr/bin/env python3
"""Reportlab-shaped drawing API backed by python-pptx native shapes.

The GEB 7365 project deck was authored for reportlab on a 960x540 canvas.
At 13.333 x 7.5 inches one canvas pixel is exactly one point, so the same
coordinates drive an editable PowerPoint with no re-layout. Everything this
emits is a real shape or a real text run, so Yasir can edit it in PowerPoint.
"""
import math
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from reportlab.pdfbase import pdfmetrics

W, H = 960.0, 540.0

DISP = ("Times New Roman", True, False)
DISPI = ("Times New Roman", False, True)
BODY = ("Times New Roman", False, False)
MONO = ("Consolas", False, False)
MONOB = ("Consolas", True, False)

# reportlab metric names, for width measurement that matches the PDF
_RL = {DISP: "Times-Bold", DISPI: "Times-Italic", BODY: "Times-Roman",
       MONO: "Courier", MONOB: "Courier-Bold"}


def rgb(hexstr):
    h = hexstr.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


class Deck:
    def __init__(self, title, author):
        self.prs = Presentation()
        self.prs.slide_width = Pt(W)
        self.prs.slide_height = Pt(H)
        cp = self.prs.core_properties
        cp.title = title
        cp.author = author
        self.blank = self.prs.slide_layouts[6]
        self.s = None
        self.n = 0

    # ---- slide lifecycle -------------------------------------------------
    def page(self, fill):
        self.s = self.prs.slides.add_slide(self.blank)
        sh = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                     self.prs.slide_width, self.prs.slide_height)
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
        sh.line.fill.background(); sh.shadow.inherit = False
        return self.s

    def notes(self, text, size=9.5, font="Calibri"):
        """Write the speaking script into the notes pane.

        Lines opening with '>>' are stage directions and are set bold so the
        eye finds them without reading the paragraph first. In PowerPoint,
        View > Notes Page puts this under the slide, one page per slide.
        """
        ns = self.s.notes_slide
        tf = ns.notes_text_frame
        tf.word_wrap = True
        # give the pane the full width of the notes page
        ph = tf._txBody.getparent()
        try:
            body = ns.placeholders[1]
            # a PowerPoint notes page is 7.5 x 10 inches, so 540 x 720 pt.
            # The slide thumbnail sits in the top 300pt; the script gets the
            # rest, inside a 54pt margin on both sides.
            body.left = Pt(54); body.top = Pt(306)
            body.width = Pt(432); body.height = Pt(392)
        except Exception:
            pass
        lines = text.rstrip().split("\n")
        tf.clear()
        for i, ln in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.space_after = Pt(0); para.space_before = Pt(0)
            para.line_spacing = 1.05
            cue = ln.lstrip().startswith(">>") or ln.strip().startswith("[")
            head = ln.strip().endswith(":") or ln.strip() in (
                "SAY:", "MESSAGE", "ROADMAP:", "HOOK")
            r = para.add_run(); r.text = ln
            f = r.font
            f.name = font; f.size = Pt(size)
            f.bold = bool(cue or head)
            f.color.rgb = rgb("#7A2E12") if cue else rgb("#101418")

    def save(self, path):
        self.prs.save(path)

    # ---- primitives ------------------------------------------------------
    def box(self, x, t, w, h, fill=None, stroke=None, lw=1, r=None, dash=None):
        shape = MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE
        sh = self.s.shapes.add_shape(shape, Pt(x), Pt(t), Pt(w), Pt(h))
        if r:
            # adj is a fraction of the shorter side
            try:
                sh.adjustments[0] = min(0.5, float(r) / max(1.0, min(w, h)))
            except Exception:
                pass
        if fill:
            sh.fill.solid(); sh.fill.fore_color.rgb = fill
        else:
            sh.fill.background()
        if stroke:
            sh.line.color.rgb = stroke; sh.line.width = Pt(lw)
            if dash:
                sh.line.dash_style = 4  # MSO_LINE_DASH_STYLE.DASH
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        st = sh._element.find(qn("p:style"))
        if st is not None:
            sh._element.remove(st)
        return sh

    def _frame(self, x, t, w, h, wrap=True):
        tb = self.s.shapes.add_textbox(Pt(x), Pt(t), Pt(w), Pt(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.auto_size = None
        return tb, tf

    def _run(self, p, s, font, size, color, track=0):
        r = p.add_run(); r.text = s
        f = r.font
        f.name = font[0]; f.size = Pt(size); f.bold = font[1]; f.italic = font[2]
        f.color.rgb = color if color is not None else rgb("#14171C")
        rPr = r._r.get_or_add_rPr()
        if track:
            rPr.set("spc", str(int(round(track * 100))))
        # force the same face for latin/complex-script slots
        for tag in ("a:latin", "a:cs"):
            e = rPr.find(qn(tag))
            if e is None:
                e = rPr.makeelement(qn(tag), {}); rPr.append(e)
            e.set("typeface", font[0])
        return r

    def txt(self, s, x, t, font=BODY, size=15, color=None, align="l", track=0):
        """One line, positioned so its cap-top lands where reportlab put it."""
        wd = self.width(s, font, size) + track * max(0, len(s) - 1) + max(4.0, size * 0.8)
        left = x if align == "l" else (x - wd / 2 if align == "c" else x - wd)
        tb, tf = self._frame(left, t - size * 0.26, wd, size * 1.5, wrap=False)
        p = tf.paragraphs[0]
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                       "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = 1.0; p.space_after = Pt(0); p.space_before = Pt(0)
        self._run(p, s, font, size, color, track)
        return tb

    def para(self, s, x, t, w, font=BODY, size=15, color=None, lead=None, align="l"):
        lead = lead or size * 1.4
        lines = self.wrap(s, w, font, size)
        tb, tf = self._frame(x, t - size * 0.26, w + 1, lead * len(lines) + size * 0.6)
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER if align == "c" else PP_ALIGN.LEFT
            p.line_spacing = Pt(lead); p.space_after = Pt(0); p.space_before = Pt(0)
            self._run(p, ln, font, size, color)
        return lead * len(lines)

    def rule(self, t, x=54, w=W - 108, color=None, lw=1):
        ln = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                         Pt(x), Pt(t), Pt(x + w), Pt(t))
        ln.line.color.rgb = color or rgb("#D8D4CB"); ln.line.width = Pt(lw)
        ln.shadow.inherit = False
        return ln

    def arrow(self, x1, t1, x2, t2, color=None, lw=1.4, dash=None, head=5):
        ln = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                         Pt(x1), Pt(t1), Pt(x2), Pt(t2))
        ln.line.color.rgb = color or rgb("#9AA4AE"); ln.line.width = Pt(lw)
        ln.shadow.inherit = False
        if dash:
            ln.line.dash_style = 4
        lnEl = ln.line._get_or_add_ln()
        tail = lnEl.makeelement(qn("a:tailEnd"),
                                {"type": "triangle", "w": "med", "len": "med"})
        lnEl.append(tail)
        return ln

    def image(self, path, x, t, w, h):
        return self.s.shapes.add_picture(path, Pt(x), Pt(t), Pt(w), Pt(h))

    # ---- metrics ---------------------------------------------------------
    def width(self, s, font, size):
        return pdfmetrics.stringWidth(s, _RL[font], size)

    def wrap(self, s, w, font, size):
        lines, line = [], ""
        for word in s.split():
            cand = (line + " " + word).strip()
            if self.width(cand, font, size) <= w:
                line = cand
            else:
                if line:
                    lines.append(line)
                line = word
        if line:
            lines.append(line)
        return lines or [""]
