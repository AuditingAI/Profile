import sys, docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
src, out, seal, mark = sys.argv[1:5]
d = docx.Document(src)
s = d.sections[0]
s.different_first_page_header_footer = True
s.header_distance = Inches(0.35)
h = s.first_page_header
for p in list(h.paragraphs): p._p.getparent().remove(p._p)
t = h.add_table(rows=1, cols=3, width=Inches(6.5))
t.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl = t._tbl; borders = OxmlElement('w:tblBorders')
for e in ('top','left','bottom','right','insideH','insideV'):
    b = OxmlElement('w:'+e); b.set(qn('w:val'),'nil'); borders.append(b)
tbl.tblPr.append(borders)
c0, c1, c2 = t.rows[0].cells
t.autofit = False
for gc, w in zip(tbl.tblGrid.findall(qn('w:gridCol')), (0.7, 5.1, 0.7)):
    gc.set(qn('w:w'), str(int(w*1440)))
for cell, w in zip((c0, c1, c2), (0.7, 5.1, 0.7)):
    cell.width = Inches(w)
lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tbl.tblPr.append(lay)
c0.paragraphs[0].add_run().add_picture(seal, height=Inches(0.5))
p = c1.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Florida International University · College of Business · DBA"); r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0x08,0x1E,0x3F); r.bold = True
p2 = c1.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("Research and materials: github.com/AuditingAI/Profile"); r2.font.size = Pt(8); r2.font.color.rgb = RGBColor(0x5B,0x65,0x77)
for pp in (p, p2):
    pp.paragraph_format.space_after = Pt(0); pp.paragraph_format.line_spacing = 1.0
q = c2.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
q.add_run().add_picture(mark, height=Inches(0.5))
d.save(out)
